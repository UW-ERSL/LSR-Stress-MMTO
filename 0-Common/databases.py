"""
databases.py
============
The material databases shipped with the repository, and the VAE settings each was
trained with, so the network shipped beside each database (.nt) can be loaded.

A database is an .xlsx file in DATABASE_DIR: row 1 the attribute names, row 2 their
units, then one row per material.
"""
import os
from dataclasses import dataclass

DATABASE_DIR = os.path.join("1-PureStructural", "MaterialDataPureStructural")


@dataclass(slots=True)
class VAEParams:
    """VAE settings for a database that has none of its own."""
    klFactor: float = 5e-6
    learningRate: float = 2e-6
    numEpochs: int = 200000                 # max epochs
    vae_hiddenDim: int = 500
    latentDim: int = 2
    maxAttributeErrorPercent: float = 1     # stops training when every attribute is this close


# the settings each shipped network was trained with
DATABASE_VAE = {
    "BM1_3Mats_2Attr":                dict(latentDim=2, learningRate=5e-4, vae_hiddenDim=50,
                                           maxAttributeErrorPercent=0.1, klFactor=1e-4, numEpochs=30000),
    "6Materials_AF_3Attr":            dict(latentDim=2, learningRate=1e-4, vae_hiddenDim=50,
                                           maxAttributeErrorPercent=0.1, klFactor=1e-4, numEpochs=100000),
    "Beam_8Mats_3Attr_v4_spectrum2D": dict(latentDim=2, learningRate=5e-4, vae_hiddenDim=50,
                                           maxAttributeErrorPercent=0.1, klFactor=1e-4, numEpochs=100000),
    "CustomAlloys20_3D":              dict(latentDim=3, learningRate=5e-4, vae_hiddenDim=50,
                                           maxAttributeErrorPercent=0.1, klFactor=1e-4, numEpochs=100000),
}


# ---------------------------------------------------------------------- attributes
class Attribute:
    """One material attribute a database may hold. measures: what it is (stiffness,
    stress, density, ...), for units; low, high: the allowed range, each end open unless
    named in closed; units: offered in the Excel template."""

    def __init__(self, name, measures, doc, low=None, high=None, closed=(), units=(), aliases=()):
        self.name, self.measures, self.doc = name, measures, doc
        self.low, self.high, self.closed = low, high, tuple(closed)
        self.units, self.aliases = tuple(units), tuple(aliases)

    def in_range(self, v):
        if self.low is not None and (v < self.low or (v == self.low and "low" not in self.closed)):
            return False
        if self.high is not None and (v > self.high or (v == self.high and "high" not in self.closed)):
            return False
        return True

    def range_text(self):
        lo = "" if self.low is None else f"{self.low} {'<=' if 'low' in self.closed else '<'} "
        hi = "" if self.high is None else f" {'<=' if 'high' in self.closed else '<'} {self.high}"
        return f"{lo}value{hi}" if (lo or hi) else "any"


_STRESS_UNITS = ("Pa", "kPa", "MPa", "GPa", "psi", "ksi")

# Every attribute a database may hold. To add one, add an entry; the check and the Excel
# template follow. A column not listed here is an error.
ATTRIBUTES = {a.name: a for a in [
    Attribute("Density", "density", "mass per unit volume", low=0,
              units=("kg/m^3", "g/cm^3", "t/mm^3", "lb/in^3")),
    Attribute("Youngs_Modulus", "stress", "Young's modulus, the stiffness", low=0, units=_STRESS_UNITS,
              aliases=("Youngs Modulus", "Young_Modulus", "E")),
    Attribute("Poissons_Ratio", "ratio", "Poisson's ratio", low=-1, high=0.5, units=("-",),
              aliases=("Poisson_Ratio", "Poissons Ratio", "nu")),
    Attribute("Yield_Strength", "stress", "yield strength, the stress at which the material yields",
              low=0, units=_STRESS_UNITS, aliases=("Yield Strength", "Yield")),
    Attribute("Cost", "cost per mass", "price per unit mass", low=0, closed=("low",),
              units=("USD/kg", "USD/lb", "USD/g")),
]}
_BY_ALIAS = {k.lower(): a.name for a in ATTRIBUTES.values() for k in (a.name,) + a.aliases}


def attribute_name(name):
    """The registry name of a database column; the name itself when it is not in the registry."""
    return _BY_ALIAS.get(str(name).strip().lower(), name)


# ---------------------------------------------------------------------- reading and checking
def _read(path):
    import warnings
    import pandas as pd
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension")
        return pd.read_excel(path, header=None)


def normalize(values, attribute_names):
    """Each column scaled to [0, 1] (log10 first for the columns named E0..E3, Y0..Y3):
    the normalized data, and each column's minimum and maximum before scaling."""
    import numpy as np
    normalized = np.zeros_like(values)
    scale_min = np.zeros(values.shape[1])
    scale_max = np.zeros(values.shape[1])
    for i, name in enumerate(attribute_names):
        col = values[:, i]
        if name in ['E0', 'E1', 'E2', 'E3', 'Y0', 'Y1', 'Y2', 'Y3']:
            col = np.log10(col)
        scale_min[i] = col.min()
        scale_max[i] = col.max()
        normalized[:, i] = (col - scale_min[i]) / (scale_max[i] - scale_min[i] + 1e-12)
    return normalized, scale_min, scale_max


def participation_ratio(X):
    """d_eff = (sum of PCA eigenvalues)^2 / sum of their squares, with the eigenvalues
    and the explained variance ratios. 0 when every eigenvalue is zero."""
    import numpy as np
    from sklearn.decomposition import PCA
    pca = PCA().fit(X)
    lambdas = np.asarray(pca.explained_variance_, dtype=float)
    if np.all(lambdas <= 1e-16):
        return 0.0, lambdas, np.zeros_like(lambdas)
    return (float((np.sum(lambdas) ** 2) / np.sum(lambdas ** 2)), lambdas,
            np.asarray(pca.explained_variance_ratio_, dtype=float))


def database_attributes(path):
    """The registry names of a database's columns, in order."""
    return [attribute_name(n) for n in _read(path).iloc[0, 1:].tolist()]


def recommended_latent_dim(path):
    """(d_eff, n_z): the participation ratio of the normalized database, and the latent
    dimension recommended for it, the smallest integer at or above d_eff (at least 1)."""
    import math
    import numpy as np
    df = _read(path)
    names = [attribute_name(n) for n in df.iloc[0, 1:].tolist()]
    values = df.iloc[2:, 1:].to_numpy(dtype=float)
    d_eff = participation_ratio(normalize(values, names)[0])[0]
    return d_eff, max(1, math.ceil(d_eff))


def check_database(path):
    """(errors, warnings) for a database file. Units are not checked."""
    import math
    import numpy as np
    from openpyxl.utils import get_column_letter
    errors, warnings_ = [], []
    if not os.path.isfile(path):
        return [f"{path} does not exist"], []
    try:
        df = _read(path)
    except Exception as e:                                          # noqa: BLE001
        return [f"{path} could not be read as an Excel file: {e}"], []

    def cell(r, c):
        return f"{get_column_letter(c + 1)}{r + 1}"

    def blank(v):
        return v is None or (isinstance(v, float) and math.isnan(v)) or str(v).strip() == ""

    if df.shape[1] < 2 or df.shape[0] < 4:
        return [f"the first sheet needs attribute names in row 1, units in row 2 and at least two "
                f"materials below, with the material names in column A; it has {df.shape[0]} rows "
                f"and {df.shape[1]} columns"], []

    names = {}
    for c in range(1, df.shape[1]):
        raw = df.iat[0, c]
        if blank(raw):
            errors.append(f"{cell(0, c)}: no attribute name")
            continue
        name = attribute_name(raw)
        if name not in ATTRIBUTES:
            errors.append(f"{cell(0, c)}: {raw!r} is not a known attribute ({', '.join(ATTRIBUTES)}); "
                          f"to use a new one, add it to ATTRIBUTES in 0-Common/databases.py")
            continue
        if name in names.values():
            errors.append(f"{cell(0, c)}: {name} appears twice")
            continue
        names[c] = name

    materials = {}
    for r in range(2, df.shape[0]):
        m = df.iat[r, 0]
        if blank(m):
            errors.append(f"{cell(r, 0)}: no material name")
        elif str(m) in materials.values():
            errors.append(f"{cell(r, 0)}: material {m} appears twice")
        else:
            materials[r] = str(m)

    for c, name in names.items():
        a, column = ATTRIBUTES[name], []
        for r in range(2, df.shape[0]):
            v = df.iat[r, c]
            where = f"{cell(r, c)} ({materials.get(r, 'row ' + str(r + 1))}, {name})"
            try:
                x = float(v)
            except (TypeError, ValueError):
                errors.append(f"{where}: {v!r} is not a number")
                continue
            if not np.isfinite(x):
                errors.append(f"{where}: no value" if blank(v) else f"{where}: {v!r} is not a finite number")
                continue
            if not a.in_range(x):
                errors.append(f"{where}: {x:g} is outside the allowed range, {a.range_text()}")
            column.append(x)
        if len(column) == df.shape[0] - 2 and len(set(column)) == 1:
            warnings_.append(f"{name} is {column[0]:g} for every material, so it tells the VAE nothing "
                             f"about the materials")
    return errors, warnings_


# ---------------------------------------------------------------------- the Excel template
def write_template(path, attributes, n_materials=5):
    """An empty database to fill in: the attributes in row 1, a units list under each in
    row 2, numbers only below, and a Reference sheet describing every known attribute."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    names = [attribute_name(a) for a in attributes]
    unknown = [a for a, n in zip(attributes, names) if n not in ATTRIBUTES]
    if unknown:
        raise ValueError(f"not known attributes: {', '.join(unknown)}; known: {', '.join(ATTRIBUTES)}")
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    ws["A1"], ws["A2"] = "Attributes", "Material/Units"
    last = 2 + n_materials
    for c, name in enumerate(names, start=2):
        col, a = get_column_letter(c), ATTRIBUTES[name]
        ws[f"{col}1"], ws[f"{col}2"] = name, a.units[0]
        units = DataValidation(type="list", formula1='"' + ",".join(a.units) + '"', allow_blank=False)
        units.error = f"Choose one of: {', '.join(a.units)}"
        ws.add_data_validation(units)
        units.add(f"{col}2")
        number = DataValidation(type="decimal", operator="between",
                                formula1=str(-1e300 if a.low is None else a.low),
                                formula2=str(1e300 if a.high is None else a.high))
        number.error = f"{name}: a number, {a.range_text()}"
        ws.add_data_validation(number)
        number.add(f"{col}3:{col}{last}")
        ws.column_dimensions[col].width = max(14, len(name) + 2)
    for r in range(3, last + 1):
        ws[f"A{r}"] = f"Material{r - 2}"
    ws.column_dimensions["A"].width = 16

    ref = wb.create_sheet("Reference")
    ref.append(["Attribute", "Meaning", "Allowed range", "Units"])
    for a in ATTRIBUTES.values():
        ref.append([a.name, a.doc, a.range_text(), ", ".join(a.units)])
    for col, w in zip("ABCD", (18, 52, 22, 36)):
        ref.column_dimensions[col].width = w
    wb.save(path)
    return path
