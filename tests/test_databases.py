"""
test_databases.py
=================
Material databases: the attribute registry, the check of a database file, the latent
dimension recommended for it, the Excel template, and what a Project does with them.
Small files written on the fly; a few seconds.

    python tests/test_databases.py        (or: python -m pytest tests)
"""
import os
import shutil
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from utils import project_paths  # noqa: E402

project_paths.bootstrap()

from openpyxl import Workbook, load_workbook  # noqa: E402
from databases import (ATTRIBUTES, DATABASE_DIR, DATABASE_VAE, VAEParams, attribute_name,  # noqa: E402
                       check_database, recommended_latent_dim, write_template)  # type: ignore

SHIPPED = os.path.join(ROOT, DATABASE_DIR)
GOOD = [["Attributes", "Density", "Youngs_Modulus", "Yield_Strength"],
        ["Material/Units", "kg/m^3", "Pa", "Pa"],
        ["A", 1.0, 100.0, 100.0],
        ["B", 1.2, 160.0, 140.0],
        ["C", 0.8, 70.0, 90.0]]


def write(path, rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(path)
    return path


def errors_of(rows):
    with tempfile.TemporaryDirectory() as d:
        return check_database(write(os.path.join(d, "db.xlsx"), rows))


def test_shipped_databases():
    """The four databases pass with no errors and no warnings, and the recommended latent
    dimension is the one each network was trained with."""
    from materialEncoder import MaterialEncoder  # type: ignore
    import contextlib
    import io
    for name, vae in DATABASE_VAE.items():
        path = os.path.join(SHIPPED, f"{name}.xlsx")
        assert check_database(path) == ([], []), name
        d_eff, n_z = recommended_latent_dim(path)
        assert n_z == vae["latentDim"], (name, d_eff, n_z)
        enc = MaterialEncoder(VAEParams())
        with contextlib.redirect_stdout(io.StringIO()):
            enc.readExcel(path)
        assert abs(enc.intrinsic_dim - d_eff) < 1e-5, (name, enc.intrinsic_dim, d_eff)


def test_errors_name_the_cell():
    assert errors_of(GOOD) == ([], [])
    cases = [
        (lambda r: r[0].__setitem__(2, "Youngs_Modulous"), "C1: 'Youngs_Modulous' is not a known attribute"),
        (lambda r: r[0].__setitem__(3, "E"), "D1: Youngs_Modulus appears twice"),
        (lambda r: r[0].__setitem__(1, None), "B1: no attribute name"),
        (lambda r: r[3].__setitem__(2, -5.0), "C4 (B, Youngs_Modulus): -5 is outside the allowed range"),
        (lambda r: r[3].__setitem__(2, "stiff"), "C4 (B, Youngs_Modulus): 'stiff' is not a number"),
        (lambda r: r[4].__setitem__(1, None), "B5 (C, Density): no value"),
        (lambda r: r[4].__setitem__(0, "A"), "A5: material A appears twice"),
        (lambda r: r[4].__setitem__(0, None), "A5: no material name"),
    ]
    for change, says in cases:
        rows = [list(r) for r in GOOD]
        change(rows)
        errors, _ = errors_of(rows)
        assert any(says in e for e in errors), (says, errors)
    errors, _ = errors_of(GOOD[:3])
    assert "at least two materials" in errors[0]
    errors, _ = errors_of([["Attributes", "Poissons_Ratio"], ["Material/Units", "-"], ["A", 0.3], ["B", 0.6]])
    assert "-1 < value < 0.5" in errors[0], errors
    assert "ATTRIBUTES in 0-Common/databases.py" in errors_of([["Attributes", "Colour"], ["u", "-"],
                                                               ["A", 1], ["B", 2]])[0][0]


def test_warnings_and_aliases():
    rows = [["Attributes", "youngs modulus", "Poisson_Ratio"], ["Material/Units", "Pa", "-"],
            ["A", 1.0, 0.3], ["B", 2.0, 0.3]]
    errors, warnings = errors_of(rows)
    assert errors == [] and len(warnings) == 1 and "Poissons_Ratio is 0.3 for every material" in warnings[0]
    assert attribute_name("Poisson_Ratio") == "Poissons_Ratio" and attribute_name("Colour") == "Colour"


def test_template():
    with tempfile.TemporaryDirectory() as d:
        path = write_template(os.path.join(d, "t.xlsx"), ["Density", "Youngs_Modulus", "Yield"], n_materials=3)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Main", "Reference"]
        main = wb["Main"]
        assert [c.value for c in main[1]] == ["Attributes", "Density", "Youngs_Modulus", "Yield_Strength"]
        lists = [v for v in main.data_validations.dataValidation if v.type == "list"]
        assert any("MPa" in v.formula1 and "C2" in str(v.sqref) for v in lists)
        assert wb["Reference"].max_row == 1 + len(ATTRIBUTES)
        errors, _ = check_database(path)                       # empty: every value is missing
        assert any("no value" in e for e in errors)
        for r, values in enumerate([(1.0, 100.0, 100.0), (1.2, 160.0, 140.0), (0.8, 70.0, 90.0)], start=3):
            for c, v in enumerate(values, start=2):
                main.cell(row=r, column=c, value=v)
        wb.save(path)
        assert check_database(path) == ([], [])
        try:
            write_template(os.path.join(d, "u.xlsx"), ["Colour"])
        except ValueError as e:
            assert "not known attributes: Colour" in str(e)
        else:
            raise AssertionError("an unknown attribute was accepted")


def test_project_uses_the_check():
    from project import Project  # type: ignore
    ubracket = os.path.join(ROOT, "0-Common", "Models", "UBracket")
    with tempfile.TemporaryDirectory() as d:
        shutil.copy(os.path.join(ubracket, "UBracket.STL"), d)
        p = Project("u", stl="UBracket.STL", base_dir=d)
        p.add_selection("bottom", __import__("selections").face("ymin"))
        p.add_support("bottom")
        p.add_force("bottom", total=[1, 0, 0])
        p.objective("FailureFactorPNorm()")
        p.constraint("Mass()", "<=", 1.0)

        p.database = "BM1_3Mats_2Attr"                         # no Yield_Strength, no Density
        errors = "\n".join(p.validate())
        assert "FailureFactorPNorm needs a Yield_Strength column" in errors, errors
        assert "Mass needs a Density column" in errors, errors

        bad = [list(r) for r in GOOD]
        bad[3][2] = -1.0
        p.database = write(os.path.join(d, "bad.xlsx"), bad)
        assert any("bad.xlsx: C4 (B, Youngs_Modulus)" in e for e in p.validate())

        p.database = write(os.path.join(d, "good.xlsx"), GOOD)
        assert p.validate() == []
        p.complete_settings()
        n_z = recommended_latent_dim(p.database)[1]
        assert p.settings["vae"]["latentDim"] == n_z                # a new database starts from the recommendation
        assert p.warnings() == []
        p.settings["vae"]["latentDim"] = n_z + 2
        assert any("is recommended" in w for w in p.warnings())


if __name__ == "__main__":
    failed = 0
    for name, fn in [(k, v) for k, v in list(globals().items()) if k.startswith("test_")]:
        try:
            fn()
            print(f"ok   {name}")
        except Exception as e:                                     # noqa: BLE001
            failed += 1
            print(f"FAIL {name}: {type(e).__name__}: {e}")
    print("all passed" if not failed else f"{failed} failed")
    sys.exit(1 if failed else 0)
